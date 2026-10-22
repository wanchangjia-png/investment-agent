#!/usr/bin/env python3
"""统一配置：路径、端口、Excel 配置表、交易费率、LLM 参数。"""
import os
import threading
from datetime import datetime, timedelta, timezone
from pathlib import Path

import openpyxl

BEIJING_TZ = timezone(timedelta(hours=8))
DATA_DIR = Path(os.path.dirname(os.path.abspath(__file__)))
EXCEL_PATH = Path(os.environ.get("PORTFOLIO_XLSX", DATA_DIR / "portfolio.xlsx"))
BACKUP_DIR = Path(os.environ.get("PORTFOLIO_BACKUP_DIR", DATA_DIR / "backups"))
HOST = os.environ.get("HOST", "0.0.0.0")
PORT = int(os.environ.get("PORT", os.environ.get("RAILWAY_PORT", 5050)))

CONFIG_SHEET = "配置"
BACKUP_KEEP_DAYS = 14

LLM_ENDPOINTS = {
    "deepseek": "https://api.deepseek.com/v1/chat/completions",
    "openai": "https://api.openai.com/v1/chat/completions",
}

TRADE_FEE_DEFAULTS = {
    "commission_rate": 0.00025,   # 佣金 万 2.5
    "commission_min": 5.0,        # 单笔最低佣金
    "stamp_rate": 0.0005,         # 印花税 0.05%，仅卖出股票
    "transfer_rate": 0.00001,     # 过户费 0.001%
}

EXCEL_LOCK = threading.RLock()
_fee_cfg = None
_fee_cfg_mtime = 0.0


def now():
    return datetime.now(BEIJING_TZ)


def save_workbook(wb, path=None):
    """原子保存：先写临时文件再替换，网页不会读到写了一半的 Excel"""
    path = Path(path or EXCEL_PATH)
    tmp = path.with_name(f".{path.stem}.tmp{path.suffix}")
    with EXCEL_LOCK:
        try:
            wb.save(tmp)
            os.replace(tmp, path)
        finally:
            if tmp.exists():
                try:
                    tmp.unlink()
                except OSError:
                    pass


def read_config_sheet():
    cfg = {}
    if not EXCEL_PATH.exists():
        return cfg
    wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True)
    try:
        if CONFIG_SHEET in wb.sheetnames:
            for row in wb[CONFIG_SHEET].iter_rows(min_row=2, max_col=2, values_only=True):
                if row[0] and row[1] not in (None, ""):
                    cfg[str(row[0]).strip()] = str(row[1]).strip()
    finally:
        wb.close()
    return cfg


def write_config_values(values):
    wb = openpyxl.load_workbook(EXCEL_PATH)
    if CONFIG_SHEET not in wb.sheetnames:
        wb.create_sheet(CONFIG_SHEET)
    ws = wb[CONFIG_SHEET]
    rows = {}
    for row in ws.iter_rows(min_row=2, max_col=2):
        if row[0].value:
            rows[str(row[0].value).strip()] = row[0].row
    for key, value in values.items():
        if key in rows:
            ws.cell(rows[key], 2, value)
        else:
            ws.append([key, value])
    save_workbook(wb)
    wb.close()


def load_llm_config():
    cfg = {
        "provider": os.environ.get("LLM_PROVIDER", "deepseek"),
        "api_key": os.environ.get("DEEPSEEK_API_KEY", os.environ.get("LLM_API_KEY", "")),
        "model": os.environ.get("LLM_MODEL", "deepseek-chat"),
    }
    try:
        sheet = read_config_sheet()
    except Exception:
        return cfg
    for key in ("provider", "api_key", "model"):
        if sheet.get(key):
            cfg[key] = sheet[key]
    return cfg


def load_fee_config(force=False):
    """费率从「配置」sheet 读取；按文件 mtime 缓存，改完立即生效"""
    global _fee_cfg, _fee_cfg_mtime
    try:
        mtime = EXCEL_PATH.stat().st_mtime
    except OSError:
        return dict(TRADE_FEE_DEFAULTS)
    if not force and _fee_cfg is not None and mtime == _fee_cfg_mtime:
        return dict(_fee_cfg)
    cfg = dict(TRADE_FEE_DEFAULTS)
    try:
        sheet = read_config_sheet()
    except Exception:
        sheet = {}
    for key in cfg:
        if key not in sheet:
            continue
        try:
            value = float(sheet[key])
        except (TypeError, ValueError):
            continue
        if value >= 0:
            cfg[key] = value
    _fee_cfg, _fee_cfg_mtime = cfg, mtime
    return dict(cfg)


def update_fee_config(values):
    cfg = load_fee_config(force=True)
    for key, value in values.items():
        if key in cfg:
            cfg[key] = float(value)
    write_config_values(cfg)
    load_fee_config(force=True)
    return cfg

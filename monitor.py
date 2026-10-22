#!/usr/bin/env python3
"""持仓日历、公告监控、卖出资金去向、不动日。

公司事件和公告来自东方财富公开接口，只在内存里缓存，不落盘。
手工维护的内容写进 portfolio.xlsx 的独立 sheet。
"""
import json
import threading
import time
import urllib.parse
import urllib.request
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta, timezone

import openpyxl

import portfolio_agent as agent

BEIJING_TZ = timezone(timedelta(hours=8))
UA = "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36"

_cache = {}
_cache_lock = threading.Lock()
CACHE_TTL = 1800
ANNOUNCE_TTL = 3600

CAL_SHEET = "持仓日历"
CAL_HEADERS = ["日期", "代码", "名称", "类型", "说明", "来源"]
SELL_SHEET = "卖出计划"
SELL_HEADERS = ["卖出日期", "代码", "名称", "卖出金额", "用途", "目标标的",
                "计划执行日", "状态", "备注", "登记时间"]
QUIET_SHEET = "不动日"
QUIET_HEADERS = ["日期", "原因", "记录时间"]

SELL_PURPOSES = ["暂存现金", "再买入", "转出", "还贷", "其他"]
RISK_WORDS = {
    "减持": "减持", "质押": "质押", "立案": "立案", "问询": "问询",
    "处罚": "处罚", "退市": "退市风险", "风险警示": "风险警示",
    "预亏": "业绩预亏", "亏损": "业绩亏损", "诉讼": "诉讼",
    "冻结": "冻结", "终止": "终止", "担保": "担保", "更正": "更正",
}


def _now():
    return datetime.now(BEIJING_TZ)


def _today():
    return _now().strftime("%Y-%m-%d")


def _cached(key, ttl, loader, force=False):
    now = time.time()
    with _cache_lock:
        hit = _cache.get(key)
        if not force and hit and now - hit[0] < ttl:
            return hit[1]
    data = loader()
    with _cache_lock:
        _cache[key] = (now, data)
    return data


def _http_json(url, timeout=12):
    req = urllib.request.Request(url, headers={"User-Agent": UA,
                                               "Referer": "https://data.eastmoney.com/"})
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        return json.loads(resp.read().decode("utf-8", "ignore"))


def _datacenter(report, code, sort_column, size=10):
    qs = urllib.parse.urlencode({
        "reportName": report, "columns": "ALL",
        "filter": f'(SECURITY_CODE="{code}")',
        "pageNumber": 1, "pageSize": size,
        "sortColumns": sort_column, "sortTypes": -1,
        "source": "WEB", "client": "WEB",
    })
    data = _http_json("https://datacenter-web.eastmoney.com/api/data/v1/get?" + qs)
    return ((data.get("result") or {}).get("data")) or []


def held_stocks():
    return [(str(h.get("代码") or "").strip(), h["名称"])
            for h in agent.load_portfolio()
            if h["类别"] == "股票" and (h.get("数量") or 0) > 0 and h.get("代码")]


def _date(value):
    return str(value or "")[:10]


def _sheet(wb, name, headers, widths):
    if name in wb.sheetnames:
        return wb[name]
    ws = wb.create_sheet(name)
    ws.append(headers)
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    return ws


# ---------- 手工日历 ----------

def _cal_sheet(wb):
    return _sheet(wb, CAL_SHEET, CAL_HEADERS,
                  {"A": 12, "B": 10, "C": 14, "D": 12, "E": 34, "F": 10})


def load_manual_events():
    try:
        wb = openpyxl.load_workbook(agent.EXCEL_PATH, read_only=True)
        if CAL_SHEET not in wb.sheetnames:
            wb.close()
            return []
        ws = wb[CAL_SHEET]
        rows = []
        for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not row[0]:
                continue
            rows.append({
                "row": i, "date": _date(row[0]), "code": str(row[1] or ""),
                "name": str(row[2] or ""), "type": str(row[3] or "其他"),
                "detail": str(row[4] or ""), "source": "手工",
            })
        wb.close()
        return rows
    except Exception as e:
        print(f"⚠️ 读取持仓日历失败: {e}")
        return []


def add_manual_event(data):
    date = _date(data.get("date"))
    if not date:
        raise ValueError("请选择日期")
    kind = str(data.get("type") or "其他").strip()
    detail = str(data.get("detail") or "").strip()
    code = str(data.get("code") or "").strip()
    name = str(data.get("name") or "").strip()
    if not name and code:
        name = next((n for c, n in held_stocks() if c == code), "")
    if not name and not detail:
        raise ValueError("请填写名称或说明")
    wb = openpyxl.load_workbook(agent.EXCEL_PATH)
    _cal_sheet(wb).append([date, code, name, kind, detail, "手工"])
    agent.save_workbook(wb)
    wb.close()
    return {"date": date, "code": code, "name": name, "type": kind, "detail": detail}


def delete_manual_event(row):
    wb = openpyxl.load_workbook(agent.EXCEL_PATH)
    ws = _cal_sheet(wb)
    if row < 2 or row > ws.max_row:
        raise ValueError("行号超出范围")
    ws.delete_rows(row)
    agent.save_workbook(wb)
    wb.close()
    return True


# ---------- 公司事件 ----------

def _events_for(code, name):
    events = []
    today = datetime.strptime(_today(), "%Y-%m-%d").date()
    horizon = today + timedelta(days=120)

    try:
        for r in _datacenter("RPT_PUBLIC_BS_APPOIN", code, "APPOINT_PUBLISH_DATE", 4):
            d = _date(r.get("APPOINT_PUBLISH_DATE"))
            if not d or r.get("IS_PUBLISH") == "1":
                continue
            dt = datetime.strptime(d, "%Y-%m-%d").date()
            if today - timedelta(days=3) <= dt <= horizon:
                events.append({"date": d, "code": code, "name": name, "type": "财报",
                               "detail": r.get("REPORT_TYPE_NAME") or "定期报告",
                               "source": "东方财富"})
    except Exception:
        pass

    try:
        for r in _datacenter("RPT_LIFT_STAGE", code, "FREE_DATE", 6):
            d = _date(r.get("FREE_DATE"))
            if not d:
                continue
            dt = datetime.strptime(d, "%Y-%m-%d").date()
            if not (today <= dt <= horizon):
                continue
            shares = r.get("CURRENT_FREE_SHARES")
            cap = r.get("LIFT_MARKET_CAP")
            detail = r.get("FREE_SHARES_TYPE") or "限售解禁"
            if shares is not None:
                detail += f"，解禁 {float(shares):,.2f} 万股"
            if cap is not None:
                detail += f"，市值约 {float(cap):,.0f} 万元"
            events.append({"date": d, "code": code, "name": name, "type": "解禁",
                           "detail": detail, "source": "东方财富"})
    except Exception:
        pass

    try:
        for r in _datacenter("RPT_SHAREBONUS_DET", code, "REPORT_DATE", 3):
            d = _date(r.get("EX_DIVIDEND_DATE") or r.get("EQUITY_RECORD_DATE"))
            if not d:
                continue
            dt = datetime.strptime(d, "%Y-%m-%d").date()
            if not (today - timedelta(days=14) <= dt <= horizon):
                continue
            detail = r.get("IMPL_PLAN_PROFILE") or "分红方案"
            events.append({"date": d, "code": code, "name": name, "type": "除权除息",
                           "detail": detail, "source": "东方财富"})
    except Exception:
        pass
    return events


def company_events(force=False):
    def loader():
        stocks = held_stocks()
        events = []
        if stocks:
            with ThreadPoolExecutor(max_workers=min(8, len(stocks))) as pool:
                for chunk in pool.map(lambda s: _events_for(*s), stocks):
                    events.extend(chunk)
        return {"events": events, "fetched_at": _now().strftime("%Y-%m-%d %H:%M"),
                "stocks": len(stocks)}

    data = dict(_cached("events", CACHE_TTL, loader, force))
    # 手工事件每次现读，增删后不用重抓网络
    events = data.get("events", []) + load_manual_events()
    events.sort(key=lambda e: (e["date"], e.get("type", "")))
    data["events"] = events
    return data


# ---------- 公告 ----------

def _announcements_for(code, name, limit=6):
    qs = urllib.parse.urlencode({
        "sr": -1, "page_size": limit, "page_index": 1, "ann_type": "A",
        "client_source": "web", "stock_list": code, "f_node": 0, "s_node": 0,
    })
    data = _http_json("https://np-anotice-stock.eastmoney.com/api/security/ann?" + qs)
    items = []
    cutoff = (datetime.strptime(_today(), "%Y-%m-%d").date() - timedelta(days=30))
    for a in ((data.get("data") or {}).get("list") or [])[:limit]:
        d = _date(a.get("notice_date")) or _date(a.get("display_time"))
        if not d:
            continue
        try:
            if datetime.strptime(d, "%Y-%m-%d").date() < cutoff:
                continue
        except ValueError:
            pass
        title = a.get("title_ch") or a.get("title") or ""
        tags = [label for word, label in RISK_WORDS.items() if word in title]
        items.append({
            "date": d, "code": code, "name": name, "title": title,
            "url": f"https://data.eastmoney.com/notices/detail/{code}/{a.get('art_code', '')}.html",
            "risk": tags,
        })
    return items


def announcements(force=False):
    def loader():
        stocks = held_stocks()
        items = []
        if stocks:
            with ThreadPoolExecutor(max_workers=min(8, len(stocks))) as pool:
                for chunk in pool.map(lambda s: _announcements_for(*s), stocks):
                    items.extend(chunk)
        items.sort(key=lambda x: x["date"], reverse=True)
        return {
            "items": items,
            "risk_count": sum(1 for i in items if i["risk"]),
            "fetched_at": _now().strftime("%Y-%m-%d %H:%M"),
            "stocks": len(stocks),
        }

    return _cached("announcements", ANNOUNCE_TTL, loader, force)


# ---------- 卖出资金去向 ----------

def _sell_sheet(wb):
    return _sheet(wb, SELL_SHEET, SELL_HEADERS,
                  {"A": 12, "B": 10, "C": 14, "D": 12, "E": 12, "F": 14,
                   "G": 12, "H": 10, "I": 24, "J": 18})


def load_sell_plans():
    try:
        wb = openpyxl.load_workbook(agent.EXCEL_PATH, read_only=True)
        if SELL_SHEET not in wb.sheetnames:
            wb.close()
            return []
        ws = wb[SELL_SHEET]
        rows = []
        for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not row[0]:
                continue
            rows.append({
                "row": i, "date": _date(row[0]), "code": str(row[1] or ""),
                "name": str(row[2] or ""), "amount": float(row[3] or 0),
                "purpose": str(row[4] or "暂存现金"), "target": str(row[5] or ""),
                "plan_date": _date(row[6]), "status": str(row[7] or "待执行"),
                "note": str(row[8] or ""), "created_at": str(row[9] or ""),
            })
        wb.close()
        return rows
    except Exception as e:
        print(f"⚠️ 读取卖出计划失败: {e}")
        return []


def unplanned_sells(days=60):
    plans = {(p["date"], p["name"]) for p in load_sell_plans()}
    cutoff = (_now() - timedelta(days=days)).strftime("%Y-%m-%d")
    sells = []
    for t in agent.load_transactions():
        if t["direction"] != "卖出" or t["date"][:10] < cutoff:
            continue
        key = (t["date"][:10], t["name"])
        if key in plans:
            continue
        sells.append({"date": t["date"][:10], "code": t.get("code", ""),
                      "name": t["name"], "amount": round(t.get("amount") or 0, 2),
                      "qty": t.get("qty")})
    sells.sort(key=lambda x: x["date"], reverse=True)
    return sells


def sell_plan_overview():
    plans = load_sell_plans()
    pending = [p for p in plans if p["status"] == "待执行"]
    cash = round(sum((h.get("市值") or 0) for h in agent.load_portfolio()
                     if h["类别"] == "现金"), 2)
    planned_in = round(sum(p["amount"] for p in pending if p["purpose"] == "再买入"), 2)
    return {
        "plans": sorted(plans, key=lambda p: (p["status"] != "待执行", p["date"]), reverse=False),
        "unplanned": unplanned_sells(),
        "cash": cash,
        "planned_rebuy": planned_in,
        "purposes": SELL_PURPOSES,
    }


def add_sell_plan(data):
    wb = openpyxl.load_workbook(agent.EXCEL_PATH)
    _sell_sheet(wb).append([
        _date(data.get("date")) or _today(),
        str(data.get("code") or ""), str(data.get("name") or ""),
        round(float(data.get("amount") or 0), 2),
        str(data.get("purpose") or "暂存现金"),
        str(data.get("target") or ""), _date(data.get("plan_date")) or None,
        str(data.get("status") or "待执行"), str(data.get("note") or ""),
        _now().strftime("%Y-%m-%d %H:%M"),
    ])
    agent.save_workbook(wb)
    wb.close()
    return True


def update_sell_plan(data):
    row = int(data.get("row") or 0)
    if row < 2:
        raise ValueError("缺少行号")
    wb = openpyxl.load_workbook(agent.EXCEL_PATH)
    ws = _sell_sheet(wb)
    if row > ws.max_row:
        wb.close()
        raise ValueError("行号超出范围")
    if data.get("purpose") is not None:
        ws.cell(row, 5, str(data["purpose"]))
    if data.get("target") is not None:
        ws.cell(row, 6, str(data["target"]))
    if data.get("plan_date") is not None:
        ws.cell(row, 7, _date(data["plan_date"]) or None)
    if data.get("status") is not None:
        ws.cell(row, 8, str(data["status"]))
    if data.get("note") is not None:
        ws.cell(row, 9, str(data["note"]))
    agent.save_workbook(wb)
    wb.close()
    return True


def delete_sell_plan(row):
    row = int(row or 0)
    wb = openpyxl.load_workbook(agent.EXCEL_PATH)
    ws = _sell_sheet(wb)
    if row < 2 or row > ws.max_row:
        wb.close()
        raise ValueError("行号超出范围")
    ws.delete_rows(row)
    agent.save_workbook(wb)
    wb.close()
    return True


# ---------- 不动日 ----------

def _quiet_sheet(wb):
    return _sheet(wb, QUIET_SHEET, QUIET_HEADERS, {"A": 12, "B": 40, "C": 18})


def load_quiet_days():
    try:
        wb = openpyxl.load_workbook(agent.EXCEL_PATH, read_only=True)
        if QUIET_SHEET not in wb.sheetnames:
            wb.close()
            return []
        rows = []
        for i, row in enumerate(wb[QUIET_SHEET].iter_rows(min_row=2, values_only=True), start=2):
            if not row[0]:
                continue
            rows.append({"row": i, "date": _date(row[0]),
                         "reason": str(row[1] or ""), "created_at": str(row[2] or "")})
        wb.close()
        return sorted(rows, key=lambda r: r["date"], reverse=True)
    except Exception as e:
        print(f"⚠️ 读取不动日失败: {e}")
        return []


def add_quiet_day(date=None, reason=""):
    date = _date(date) or _today()
    wb = openpyxl.load_workbook(agent.EXCEL_PATH)
    ws = _quiet_sheet(wb)
    target = None
    for i, row in enumerate(ws.iter_rows(min_row=2, max_col=2, values_only=True), start=2):
        if _date(row[0]) == date:
            target = i
            break
    if target:
        ws.cell(target, 2, reason or "今天不动")
    else:
        ws.append([date, reason or "今天不动", _now().strftime("%Y-%m-%d %H:%M")])
    agent.save_workbook(wb)
    wb.close()
    return date


def delete_quiet_day(row):
    row = int(row or 0)
    wb = openpyxl.load_workbook(agent.EXCEL_PATH)
    ws = _quiet_sheet(wb)
    if row < 2 or row > ws.max_row:
        wb.close()
        raise ValueError("行号超出范围")
    ws.delete_rows(row)
    agent.save_workbook(wb)
    wb.close()
    return True


def quiet_overview():
    rows = load_quiet_days()
    today = _today()
    cutoff = (_now() - timedelta(days=30)).strftime("%Y-%m-%d")
    return {
        "today": today,
        "active": any(r["date"] == today for r in rows),
        "days": rows,
        "last30": sum(1 for r in rows if r["date"] >= cutoff),
    }

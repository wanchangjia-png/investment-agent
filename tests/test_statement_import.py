import openpyxl

import statement_import as si


def test_query_range():
    assert si._query_range({"query_range": "2026/09/22-2026/10/22"}) == (
        "2026-09-22", "2026-10-22")
    assert si._query_range({}) == (None, None)


def test_merge_opening_prefers_ledger():
    reverse = {"002475": {"name": "立讯精密", "qty": 200}}
    ledger = {"002475": [100, 53.15]}
    out = si._merge_opening(reverse, ledger, {"002475": 105.509})
    assert out["002475"] == {"qty": 100, "cost": 53.15}


def test_merge_opening_falls_back_to_statement_cost():
    out = si._merge_opening({"510760": {"name": "上证指数ETF", "qty": 7500}}, {},
                            {"510760": 1.267})
    assert out["510760"] == {"qty": 7500, "cost": 1.267}


def test_upsert_keeps_rows_outside_range():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "交易流水"
    ws.append(si.LEDGER_HEADERS)
    old = ["2026-07-24", "买入", "股票", "三花智控", "002050", 100, 35, 3500,
           "A股主账户", "对账单导入", 5, None, 35, 100, 35]
    window = ["2026-09-24", "买入", "股票", "立讯精密", "002475", 100, 52.08, 5208,
              "A股主账户", "对账单导入", 5, None, 0, 100, 52.13]
    stale = window[:]
    stale[5] = 999
    ws.append(old)
    ws.append(stale)

    si._upsert_transactions(ws, [window], "2026-09-22", "2026-10-22")

    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 2
    assert rows[0][0] == "2026-07-24"
    assert rows[1][0] == "2026-09-24"
    assert rows[1][5] == 100


def _trade(date, type_, code, name, qty, price=0.0, page=1):
    return {"date": date, "type": type_, "code": code, "name": name,
            "qty": qty, "price": price, "amount": qty * price, "balance": 0.0,
            "page": page, "fee": 0.0, "stamp": 0.0, "transfer": 0.0}


def test_opening_positions_reverses_share_adds():
    trades = [_trade("2026-07-17", "股份转入", "159327", "半导体设备", 8000)]
    summary = [{"code": "159327", "name": "半导体设备", "qty": 11300}]
    assert si.opening_positions(trades, summary)["159327"]["qty"] == 3300


def test_build_ledger_zero_cost_share_adds():
    trades = [
        _trade("2025-06-03", "红股入账", "300058", "蓝色光标", 80, page=1),
        _trade("2026-07-17", "股份转入", "159327", "半导体设备", 8000, page=2),
    ]
    opening = {"300058": {"qty": 1000, "cost": 10.0},
               "159327": {"qty": 3300, "cost": 1.1655}}
    rows, warnings, state = si.build_ledger({"trades": trades, "summary": []}, opening)

    assert not warnings
    red, split = rows
    assert red[1] == "买入" and red[7] == 0 and red[11] is None
    assert "红股入账" in red[9]
    assert red[14] == round(1000 * 10.0 / 1080, 3)
    assert "股份转入" in split[9]
    assert state["159327"] == [11300, round(3300 * 1.1655 / 11300, 3)]


def test_build_ledger_unknown_opening_cost_skips_realized():
    trades = [_trade("2024-11-06", "证券卖出", "002146", "荣盛发展", 6000, 1.21)]
    opening = {"002146": {"qty": 6000, "cost": None}}
    rows, warnings, state = si.build_ledger({"trades": trades, "summary": []}, opening)

    assert not warnings
    row = rows[0]
    assert row[11] is None
    assert "成本未知" in row[9]
    assert state["002146"] == [0, 0.0]


def test_ledger_opening_state(tmp_path):
    path = tmp_path / "book.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "交易流水"
    ws.append(si.LEDGER_HEADERS)
    ws.append(["2026-09-17", "卖出", "股票", "三花智控", "002050", 100, 34.5, 3450,
               "A股主账户", "对账单导入", 6, -320, 37.6, 800, 37.638])
    ws.append(["2026-10-19", "卖出", "股票", "三花智控", "002050", 100, 31.92, 3192,
               "A股主账户", "对账单导入", 6, -578, 37.638, 700, 37.638])
    wb.save(path)

    state = si._ledger_opening_state(str(path), "2026-09-22")
    assert state["002050"] == [800, 37.638]

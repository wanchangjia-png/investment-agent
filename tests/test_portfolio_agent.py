import datetime as dt

import openpyxl

import portfolio_agent as agent


def test_realized_total_consistency(excel):
    ledger = agent.get_ledger_realized_total()
    legacy = sum(r["pnl"] for r in agent.get_legacy_realized_records())
    total = agent.get_realized_pnl_total()
    if ledger or legacy:
        assert total == round(ledger + legacy, 2)
    else:
        assert isinstance(total, float)


def test_corporate_action_add_delete(excel):
    before = agent.get_corporate_action_total()
    agent.add_corporate_action("2099-04-01", "000000", "测试股份", "现金分红", 88.8, "pytest")
    assert agent.get_corporate_action_total() == round(before + 88.8, 2)

    record = next(a for a in agent.load_corporate_actions() if a["note"] == "pytest")
    agent.delete_corporate_action(record["row"])
    assert agent.get_corporate_action_total() == before


def test_calculate_includes_realized(excel):
    holdings = agent.load_portfolio()
    total_value, total_pnl, accounts = agent.calculate(holdings)
    assert total_value >= 0
    assert accounts
    expected = sum((h.get("市值") or 0) for h in holdings
                   if h["类别"] in ("股票", "现金", "黄金"))
    assert round(total_value, 2) == round(expected, 2)


def _fixed_time(monkeypatch, hour, minute):
    class FixedDatetime(dt.datetime):
        @classmethod
        def now(cls, tz=None):
            return dt.datetime(2099, 4, 2, hour, minute, tzinfo=tz)

    monkeypatch.setattr(agent, "datetime", FixedDatetime)


def test_intraday_snapshot_updates_same_row(excel, monkeypatch):
    _fixed_time(monkeypatch, 10, 30)
    before = len(agent.load_history())

    assert agent.save_networth_snapshot(200000, -50000, {"A股主账户": 150000}) is True
    assert len(agent.load_history()) == before + 1

    assert agent.save_networth_snapshot(201000, -49000, {"A股主账户": 151000}) is True
    history = agent.load_history()
    assert len(history) == before + 1
    assert history[-1]["总资产"] == 201000

    wb = openpyxl.load_workbook(str(excel))
    note = wb["净值历史"].cell(row=wb["净值历史"].max_row, column=8).value
    assert note.startswith("盘中记录")


def test_close_snapshot_finalizes_once(excel, monkeypatch):
    _fixed_time(monkeypatch, 15, 10)
    before = len(agent.load_history())

    assert agent.save_networth_snapshot(200000, -50000, {"A股主账户": 150000}) is True
    assert agent.save_networth_snapshot(201000, -49000, {"A股主账户": 151000}) is False
    history = agent.load_history()
    assert len(history) == before + 1
    assert history[-1]["总资产"] == 200000

    wb = openpyxl.load_workbook(str(excel))
    note = wb["净值历史"].cell(row=wb["净值历史"].max_row, column=8).value
    assert note == "收盘记录"

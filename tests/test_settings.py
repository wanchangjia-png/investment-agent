import openpyxl

import portfolio_agent as agent
import settings


def test_fee_config_roundtrip(excel):
    settings.update_fee_config({"commission_rate": 0.0003, "commission_min": 1.0})
    cfg = settings.load_fee_config(force=True)
    assert cfg["commission_rate"] == 0.0003
    assert cfg["commission_min"] == 1.0

    settings.update_fee_config({"commission_min": 5.0})
    assert settings.load_fee_config(force=True)["commission_min"] == 5.0


def test_estimate_trade_fee_uses_config(excel):
    settings.update_fee_config({"commission_rate": 0.00025, "commission_min": 5.0,
                                "stamp_rate": 0.0005, "transfer_rate": 0.00001})
    buy = agent.estimate_trade_fee("买入", 100, 10.0)
    assert buy["commission"] == 5.0        # 0.25 元不足最低佣金
    assert buy["stamp"] == 0
    assert buy["transfer"] == 0.01
    assert buy["total"] == 5.01

    sell = agent.estimate_trade_fee("卖出", 100, 10.0)
    assert sell["stamp"] == 0.5
    assert sell["total"] == 5.51


def test_large_trade_uses_rate(excel):
    settings.update_fee_config({"commission_rate": 0.00025, "commission_min": 5.0})
    fee = agent.estimate_trade_fee("买入", 10000, 10.0)
    assert fee["commission"] == 25.0       # 10 万 × 万 2.5


def test_atomic_save_leaves_no_temp_file(excel):
    wb = openpyxl.load_workbook(excel)
    wb["持仓表"].cell(1, 1)  # touch
    settings.save_workbook(wb)
    wb.close()

    assert excel.exists()
    leftovers = list(excel.parent.glob(f".{excel.stem}.tmp*"))
    assert leftovers == []
    openpyxl.load_workbook(excel).close()
